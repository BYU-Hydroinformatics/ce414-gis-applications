"""One-off (October 28, 2026): bring lab09-report-template.docx in line with the Week 10 cut of
Step 9 (two runs, no spline, no cells-below-the-plain column, 2-page write-up).
Run from the repository root: python -X utf8 tools/lab09/cut_template.py
"""
import copy
import sys

import docx

PATH = "docs/assignments/lab-09/lab09-report-template.docx"
d = docx.Document(PATH)
changed = []


def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text)
    else:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    changed.append(text[:70])


def para_with(prefix):
    hits = [p for p in d.paragraphs if p.text.strip().startswith(prefix)]
    if len(hits) != 1:
        sys.exit(f"expected one paragraph starting {prefix!r}, found {len(hits)}")
    return hits[0]


# Section 5 — Step 9
set_text(para_with("Rubric: one table with the baseline and the four Step 9 runs"),
         "Rubric: one table with the baseline and the two Step 9 runs (250 points; the wide outline), "
         "with the points kept, the outline area, the volume and the tallest cell, and the numbers match "
         "a run with your seed (5 points). Step 9 says where to read each value.")
set_text(para_with("Table 4. Sensitivity"), "Table 4. The baseline and the two Step 9 runs.")
set_text(para_with("Which choice moves the volume most"),
         "Which choice moved the volume more, the points or the outline? Give each run's difference from "
         "the baseline in cubic kilometers and as a percentage. (3 points)")
set_text(para_with("What did the spline do that IDW cannot"),
         "What did the wider outline change: the area it counts, the plain it rebuilt, or both? Argue from "
         "your table's columns. (2 points)")
set_text(para_with("Rubric: what the volume measures and what the model cannot see"),
         "Rubric: what the volume measures and what the model cannot see (2 points). Step 9, question 2: "
         "is it the volume of the lava dome? Say what part of the dome the model cannot see, and what data "
         "would let you measure it.")
set_text(para_with("Map 2. Scenario"), "Map 2. The wide-outline run from Step 9, same color scale as Map 1.")

# Table 4 (index 6): drop the 4,000 and Spline rows, rename the outline row, drop the last column.
t = d.tables[6]
assert t.rows[0].cells[0].text == "Run", t.rows[0].cells[0].text
for row in list(t.rows):
    if row.cells[0].text.strip() in ("4,000 points", "Spline"):
        row._tr.getparent().remove(row._tr)
        changed.append(f"removed row {row.cells[0].text.strip()}")
for row in t.rows:
    if row.cells[0].text.strip() == "Your own outline":
        set_text(row.cells[0].paragraphs[0], "Wide outline (+200 m)")
assert t.rows[0].cells[-1].text.strip() == "Cells below the plain"
for row in t.rows:
    tc = row.cells[-1]._tc
    tc.getparent().remove(tc)
grid = t._tbl.tblGrid
cols = grid.findall(docx.oxml.ns.qn("w:gridCol"))
grid.remove(cols[-1])
changed.append("removed column Cells below the plain")

# Rubric table (index 9)
r = d.tables[9]
for row in r.rows:
    c = row.cells[0]
    head = c.paragraphs[0].text
    if head.startswith("Write-up"):
        set_text(c.paragraphs[0], "Write-up (2 pages)")
        for p in c.paragraphs:
            if "Step 9, question 3" in p.text:
                set_text(p, p.text.replace("Step 9, question 3", "Step 9, question 2"))
    elif head.startswith("Map 2"):
        set_text(c.paragraphs[0], "Map 2 — the wide outline (full page, 8.5 × 11)")
    elif head.startswith("Sensitivity"):
        set_text(c.paragraphs[0], "Step 9 — two more runs")
        body = c.paragraphs[1:]
        lines = [
            "One table with the baseline and the two Step 9 runs (250 points; the wide outline), with the "
            "points kept, the outline area, the volume and the tallest cell, and the numbers match a run "
            "with your seed (5)",
            "Which choice moved the volume more, with each run's difference from the baseline in cubic "
            "kilometers and as a percentage (3)",
            "What the wider outline changed, the area counted or the plain rebuilt or both, argued from "
            "your table (2)",
        ]
        for p, line in zip(body, lines):
            set_text(p, line)
        for p in body[len(lines):]:
            p._p.getparent().remove(p._p)
        if len(body) < len(lines):
            sys.exit("rubric Sensitivity cell has fewer paragraphs than expected")

# Any leftover mentions
for p in d.paragraphs:
    for bad in ("spline", "four Step 9", "own outline", "2–3 pages", "cells below"):
        if bad.lower() in p.text.lower():
            print("LEFTOVER:", p.text[:120])
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for bad in ("spline", "four Step 9", "own outline", "2–3 pages", "cells below"):
                if bad.lower() in c.text.lower():
                    print("LEFTOVER CELL:", c.text[:120])

d.save(PATH)
print("\n".join(changed))

"""One-off (October 28, 2026, third pass): the Step 9 widening is personal (3 x BYU ID digits).
Run from the repository root after cut_template2.py: python -X utf8 tools/lab09/cut_template3.py
"""
import copy
import sys

import docx
from docx.text.paragraph import Paragraph

PATH = "docs/assignments/lab-09/lab09-report-template.docx"
d = docx.Document(PATH)
changed = []


def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text)
    else:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    changed.append(text[:70])


def para_with(prefix):
    hits = [p for p in d.paragraphs if p.text.strip().startswith(prefix)]
    if len(hits) != 1:
        sys.exit(f"expected one paragraph starting {prefix!r}, found {len(hits)}")
    return hits[0]


rub = para_with("Rubric: one table with the baseline and the two Step 9 runs")
set_text(rub, "Rubric: one table with the baseline and the two Step 9 runs (your wide outline; the spline base), "
              "with the points kept, the outline area, the volume, the tallest cell and the lowest cell; your "
              "digits and distance stated, and the numbers match a run at your distance (5 points). Step 9 "
              "says where to read each value.")
# a line for the digits and distance, right after the rubric line
newp = copy.deepcopy(rub._p)
rub._p.addnext(newp)
set_text(Paragraph(newp, rub._parent), "Last two digits of my BYU ID: [dd].  Outline widened by 3 × [dd] = [___] m.")

t = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == "Run")
for row in t.rows:
    if row.cells[0].text.strip().startswith("Wide outline"):
        set_text(row.cells[0].paragraphs[0], "Wide outline (+[___] m, mine)")

r = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == "Item")
for row in r.rows:
    c = row.cells[0]
    if c.paragraphs[0].text.startswith("Step 9"):
        p = c.paragraphs[1]
        set_text(p, "One table with the baseline and the two Step 9 runs (your wide outline; the spline base), "
                    "with the points kept, the outline area, the volume, the tallest cell and the lowest cell; "
                    "your digits and distance stated, and the numbers match a run at your distance (5)")

for p in d.paragraphs:
    if "+200" in p.text or "200 m" in p.text:
        print("LEFTOVER:", p.text[:120])
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            if "+200" in c.text:
                print("LEFTOVER CELL:", c.text[:120])
d.save(PATH)
print("\n".join(changed))

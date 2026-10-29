"""One-off (October 28, 2026, second pass): drop the metadata section from
lab09-report-template.docx and make Step 9 the wide outline + spline base, Map 2 the spline.
Run from the repository root after cut_template.py: python -X utf8 tools/lab09/cut_template2.py
"""
import copy
import sys

import docx
from docx.oxml.ns import qn

PATH = "docs/assignments/lab-09/lab09-report-template.docx"
d = docx.Document(PATH)
changed = []


def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text)
    else:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    changed.append(text[:70])


def para_with(prefix):
    hits = [p for p in d.paragraphs if p.text.strip().startswith(prefix)]
    if len(hits) != 1:
        sys.exit(f"expected one paragraph starting {prefix!r}, found {len(hits)}")
    return hits[0]


def remove(el):
    el.getparent().remove(el)


# 1. Drop section 2 (heading, rubric line, placeholder, caption) and its table (Table 1 metadata).
h2 = para_with("2.  The DEM and Its Metadata")
h3 = para_with("3.  The Model")
body = d.element.body
start = list(body).index(h2._p)
end = list(body).index(h3._p)
for el in list(body)[start:end]:
    remove(el)
changed.append("removed section 2 (metadata) and its table")

# 2. Renumber headings 3..10 -> 2..9.
for p in d.paragraphs:
    if p.style.name == "Heading 1":
        txt = p.text.strip()
        num, _, rest = txt.partition(".")
        if num.isdigit() and int(num) >= 3:
            set_text(p, f"{int(num) - 1}.{rest}")

# 3. Section texts.
set_text(para_with("Rubric: one table with the baseline and the two Step 9 runs"),
         "Rubric: one table with the baseline and the two Step 9 runs (the wide outline; the spline base), "
         "with the points kept, the outline area, the volume, the tallest cell and the lowest cell, and "
         "the numbers match a run with your seed (5 points). Step 9 says where to read each value.")
set_text(para_with("Which choice moved the volume more"),
         "Which choice moved the volume more, the outline or the method? Give each run's difference from "
         "the baseline in cubic kilometers and as a percentage. (2 points)")
set_text(para_with("What did the wider outline change"),
         "What did the spline do that IDW cannot? Your lowest-cell column and Map 2 show it; explain it "
         "with what you learned about the two methods in Week 9 and Lab 8, and say which base you would "
         "trust for the volume, and why. (3 points here, 2 in the write-up)")
set_text(para_with("Rubric: what the volume measures and what the model cannot see"),
         "Rubric: what the volume measures and what the model cannot see (2 points). Step 9, question 3: "
         "is it the volume of the lava dome? Say what part of the dome the model cannot see, and what data "
         "would let you measure it.")
set_text(para_with("Everything Map 1 needs except the inset"),
         "Everything Map 1 needs except the inset, in the same color scale as Map 1 plus one class of its "
         "own for cells below zero, where the spline plain rose above the ground. The title and text box "
         "say which method made the base, the volume, and how far it is from Map 1's.")
set_text(para_with("Map 2. The wide-outline run"), "Map 2. The spline base from Step 9, same color scale as Map 1.")

# 4. Step 9 table: rows Baseline / Wide outline / Spline; add a Lowest cell column.
t = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == "Run")
for row in list(t.rows):
    if row.cells[0].text.strip() == "250 points":
        remove(row._tr)
        changed.append("removed row 250 points")
last = t.rows[-1]
new = copy.deepcopy(last._tr)
last._tr.addnext(new)
rows = t.rows
set_text(rows[-1].cells[0].paragraphs[0], "Spline base")
for c in rows[-1].cells[1:]:
    set_text(c.paragraphs[0], " ")
for row in t.rows:
    src = row.cells[-1]._tc
    src.addnext(copy.deepcopy(src))
grid = t._tbl.tblGrid
cols = grid.findall(qn("w:gridCol"))
grid.append(copy.deepcopy(cols[-1]))
set_text(t.rows[0].cells[-1].paragraphs[0], "Lowest cell (m)")
for row in t.rows[1:]:
    set_text(row.cells[-1].paragraphs[0], " ")
changed.append("Step 9 table: Baseline / Wide outline / Spline base, + Lowest cell column")

# 5. Rubric table.
r = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == "Item")
for row in r.rows:
    c = row.cells[0]
    head = c.paragraphs[0].text
    if head.startswith("Write-up"):
        paras = c.paragraphs[1:]
        keep = [p for p in paras if "metadata" not in p.text.lower()]
        for p in paras:
            if "metadata" in p.text.lower():
                remove(p._p)
        for p in keep:
            if "Step 9, question 2" in p.text:
                set_text(p, p.text.replace("Step 9, question 2", "Step 9, question 3"))
            if p.text.startswith("Organized writing"):
                newp = copy.deepcopy(p._p)
                p._p.addprevious(newp)
                from docx.text.paragraph import Paragraph
                set_text(Paragraph(newp, p._parent),
                         "Which base you would trust for the volume, IDW or spline, and why (Step 9, question 2) (2)")
    elif head.startswith("Map 2"):
        set_text(c.paragraphs[0], "Map 2 — the spline base (full page, 8.5 × 11)")
        lines = [
            "Title, with the method and the volume in it or in a text box (1)",
            "Neat line, north arrow and scale bar (1)",
            "Text box with author, date, map projection, and the DEM's source and date (1)",
            "Height above the plain in the same color scale as Map 1, with a legend, plus one class for cells below zero (3)",
            "The outline over a hillshade or imagery (1)",
            "Text box says how far the volume is from Map 1's (1)",
            "Scale and legibility appropriate to the butte (2)",
        ]
        body_p = c.paragraphs[1:]
        if len(body_p) != len(lines):
            sys.exit(f"Map 2 cell has {len(body_p)} lines, expected {len(lines)}")
        for p, line in zip(body_p, lines):
            set_text(p, line)
    elif head.startswith("Step 9"):
        lines = [
            "One table with the baseline and the two Step 9 runs (the wide outline; the spline base), with the "
            "points kept, the outline area, the volume, the tallest cell and the lowest cell, and the numbers "
            "match a run with your seed (5)",
            "Which choice moved the volume more, with each run's difference from the baseline in cubic "
            "kilometers and as a percentage (2)",
            "What the spline did that IDW cannot, explained from the lowest-cell column and Map 2 (3)",
        ]
        body_p = c.paragraphs[1:]
        if len(body_p) != len(lines):
            sys.exit(f"Step 9 cell has {len(body_p)} lines, expected {len(lines)}")
        for p, line in zip(body_p, lines):
            set_text(p, line)

# 6. References line mentions both tiles; fine. Check leftovers.
for p in d.paragraphs:
    for bad in ("metadata", "wide-outline run", "wider outline change", "250 points"):
        if bad.lower() in p.text.lower():
            print("LEFTOVER:", p.text[:120])
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for bad in ("metadata", "wide-outline run", "250 points"):
                if bad.lower() in c.text.lower():
                    print("LEFTOVER CELL:", c.text[:120])

d.save(PATH)
print("\n".join(changed))

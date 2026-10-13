"""Oct 12, 2026: add the AI-proofing items to lab06-report-template.docx in place.

- "Before you submit" box: three-file upload, History and attribute-table captures, personal level.
- Section 3: a Figure 3 slot for the Geoprocessing History and attribute-table captures.
- Table 4 (range and step): a pre-filled personal-level row.
Run with the standalone Python (python-docx). Idempotent: skips anything already present.
"""
import copy
import pathlib
import sys

import docx
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = pathlib.Path(__file__).resolve().parent
DOCX = HERE.parents[1] / "docs" / "assignments" / "lab-06" / "lab06-report-template.docx"
d = docx.Document(str(DOCX))
text = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)


def para_after(anchor, txt, style):
    """Insert a new paragraph with `style` directly after paragraph `anchor`."""
    new = copy.deepcopy(anchor._p)
    for child in list(new):
        if child.tag != qn("w:pPr"):
            new.remove(child)
    anchor._p.addnext(new)
    p = docx.text.paragraph.Paragraph(new, anchor._parent)
    p.style = d.styles[style]
    p.add_run(txt)
    return p


# 1. checklist box (table 0, one cell of paragraphs)
cell = d.tables[0].rows[0].cells[0]
adds = ["Upload three files: this report as a PDF, and each map as its own PDF exported from the ArcGIS Pro layout (Share ▸ Export Layout ▸ PDF). The grader checks that the map PDFs have selectable text.",
        "Section 3 holds your Geoprocessing History capture (Analysis ▸ History, showing your runs) and the attribute table of your baseline Shorelines.",
        "Table 4 has your personal-level row: your BYU ID's last two digits, the elevation they give, the main-pool area and the number of water bodies."]
have = cell.text
for a in adds:
    if a[:40] not in have:
        cell.add_paragraph(a, style=cell.paragraphs[-1].style)

# 2. Figure 3 slot after the Figure 2 caption
paras = d.paragraphs
cap2 = next(p for p in paras if p.text.startswith("Figure 2."))
if not any(p.text.startswith("Figure 3.") for p in paras):
    rubric = para_after(cap2, "Rubric: evidence that the model ran on your machine (part of the 4 points for a working model). "
                        "Paste two screen captures: the Geoprocessing History pane (Analysis tab ▸ History) listing your "
                        "model's runs with their dates and run times, and the attribute table of your baseline Shorelines "
                        "feature class showing its rows with Elevation and AreaSqMi.", "Normal")
    para_after(rubric, "Figure 3. Geoprocessing History and the baseline Shorelines attribute table.", "Caption")

# 3. personal-level row in the range-and-step table
t = d.tables[6]
if not any("Personal level" in r.cells[0].text for r in t.rows):
    row = copy.deepcopy(t.rows[1]._tr)       # the baseline row's formatting
    t.rows[0]._tr.addnext(row)                 # directly under the header, above the baseline
r = next(r for r in t.rows if "Personal level" in r.cells[0].text)
vals = ["Personal level: BYU ID digits __, E = 3,500 + 2 × __ = ____ ft", "E", "E", "any", "1", "", "same", "",
        "shoreline length at E: ______ m"]
for c, v in zip(r.cells, vals):
    for p in c.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    c.paragraphs[0].text = v
cell = d.tables[0].rows[0].cells[0]
for p in cell.paragraphs:
    if "number of water bodies" in p.text:
        p.text = p.text.replace("the main-pool area and the number of water bodies", "the main-pool area and its shoreline length")

d.save(str(DOCX))
print("updated", DOCX.name)

#!/usr/bin/env python3
"""Build the workbook behind the "Simple Map Algebra in Excel" figure on Week 3, Part A.

    python3 tools/week03_map_algebra_xlsx.py [OUT.xlsx]

The figure on that slide used to be a scan of Bolstad's Excel 2003 screenshot. This writes the same
example in a real workbook so it can be opened in current Excel and captured: two 3x3 input blocks
and a third block that is *formulas*, not numbers, so the formula bar reads `=B3+H3` — the string
the slide asks students to type.

Cell placement is chosen for that formula: Layer A starts at B3, Layer B at H3, so the first output
cell is `=B3+H3`. Change the placement and the slide text has to change with it.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

A = [[5, 7, 2], [9, 10, 4], [8, 8, 3]]
B = [[2, 3, 1], [2, 1, 4], [0, 2, 3]]

NAVY = "FF002E5D"
EDGE = Side(style="medium", color=NAVY)
GRID = Side(style="thin", color="FFBFC9D6")


def block(ws, top, left, label, fill, values=None, formula=None):
    """Write a 3x3 block with a label above it. `values` writes numbers, `formula` a callable
    returning the formula text for (row, col) offsets."""
    ws.cell(row=top - 1, column=left, value=label).font = Font(bold=True, size=12, color=NAVY)
    for r in range(3):
        for c in range(3):
            cell = ws.cell(row=top + r, column=left + c)
            cell.value = values[r][c] if values is not None else formula(r, c)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(size=14)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.border = Border(
                top=EDGE if r == 0 else GRID, bottom=EDGE if r == 2 else GRID,
                left=EDGE if c == 0 else GRID, right=EDGE if c == 2 else GRID)


def main(out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Map algebra"

    block(ws, 3, 2, "Layer A", "FFE8F0F8", values=A)          # B3:D5
    block(ws, 3, 8, "Layer B", "FFEAF3EA", values=B)          # H3:J5
    # Same cell in, same cell out: every output cell adds the two inputs at its own position.
    block(ws, 9, 2, "A + B  (formulas)", "FFFDF1E7", formula=lambda r, c:
          f"={get_column_letter(2 + c)}{3 + r}+{get_column_letter(8 + c)}{3 + r}")

    ws["F4"] = "+"
    ws["F4"].font = Font(bold=True, size=20, color=NAVY)
    ws["F4"].alignment = Alignment(horizontal="center", vertical="center")

    for col in range(1, 12):
        ws.column_dimensions[get_column_letter(col)].width = 7.5
    for row in range(1, 13):
        ws.row_dimensions[row].height = 26
    ws.sheet_view.showGridLines = True
    ws.sheet_view.zoomScale = 160
    # Land the cursor on the first formula cell so the formula bar shows =B3+H3 on open.
    ws.sheet_view.selection[0].activeCell = "B9"
    ws.sheet_view.selection[0].sqref = "B9"

    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else
         str(Path.home() / "AppData/Local/Temp/ce414-map-algebra.xlsx"))
